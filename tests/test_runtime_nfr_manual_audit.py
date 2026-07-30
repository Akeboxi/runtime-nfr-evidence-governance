from __future__ import annotations

from pathlib import Path
import shutil
import sys
import types

import pandas as pd
import pytest
from openpyxl import load_workbook

graphbolt_stub = types.ModuleType("dgl.graphbolt")
graphbolt_stub.__file__ = "<unused-graphbolt-test-stub>"
graphbolt_stub.__getattr__ = lambda name: type("UnusedGraphBoltType", (), {})
sys.modules.setdefault("dgl.graphbolt", graphbolt_stub)

from src.data.runtime_nfr_manual_audit import (  # noqa: E402
    CONTROL_AUDIT_FIELDS,
    PAIR_AUDIT_FIELDS,
    PRIVATE_COMPUTED_JUDGMENT_FIELDS,
    PRIVATE_OUTCOME_FIELDS,
    _blinded,
    _format_workbook,
    _window_id,
    summarize_manual_audit_results,
)


def test_manual_audit_blinding_removes_every_outcome_field() -> None:
    source = pd.DataFrame(
        [
            {
                "audit_pair_id": "audit-pair-001",
                "event_id": "event-1",
                "app_id": "frontend",
                "event_breach": 1,
                "control_breach": 0,
                "control_breach_rate": 0.0,
                "event_minus_control": 1.0,
                "direction": "positive",
            }
        ]
    )
    blinded = _blinded(source, CONTROL_AUDIT_FIELDS)
    assert not set(PRIVATE_OUTCOME_FIELDS) & set(blinded.columns)
    assert not set(PRIVATE_COMPUTED_JUDGMENT_FIELDS) & set(blinded.columns)
    assert set(CONTROL_AUDIT_FIELDS).issubset(blinded.columns)
    assert _window_id("frontend", pd.Timestamp("2025-01-01").to_pydatetime()) == (
        _window_id("frontend", pd.Timestamp("2025-01-01").to_pydatetime())
    )


def test_reviewer_workbook_has_only_blinded_sheets_and_validations(
    tmp_path: Path,
) -> None:
    pair = _blinded(
        pd.DataFrame(
            [
                {
                    "audit_pair_id": "audit-pair-001",
                    "event_id": "event-1",
                    "app_id": "frontend",
                    "event_breach": 1,
                    "direction": "positive",
                }
            ]
        ),
        PAIR_AUDIT_FIELDS,
    )
    control = _blinded(
        pd.DataFrame(
            [
                {
                    "audit_pair_id": "audit-pair-001",
                    "control_slot": 1,
                    "event_id": "event-1",
                    "app_id": "frontend",
                    "control_breach": 0,
                }
            ]
        ),
        CONTROL_AUDIT_FIELDS,
    )
    path = tmp_path / "reviewer.xlsx"
    _format_workbook(
        path,
        pair_frame=pair,
        control_frame=control,
        include_private=False,
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["说明", "组级核查", "逐对照核查", "字段字典"]
    headers = {
        cell.value
        for sheet_name in ("组级核查", "逐对照核查")
        for cell in workbook[sheet_name][1]
    }
    assert not set(PRIVATE_OUTCOME_FIELDS) & headers
    assert workbook["组级核查"].data_validations.count > 0
    assert workbook["逐对照核查"].data_validations.count > 0


def _signed_audit_fixture(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    pair = pd.DataFrame(
        [
            {
                "audit_pair_id": f"audit-pair-{index:03d}",
                "event_id": f"event-{index}",
                "app_id": f"app-{index % 3}",
                **{field: "" for field in PAIR_AUDIT_FIELDS},
            }
            for index in range(1, 21)
        ]
    )
    control = pd.DataFrame(
        [
            {
                "audit_pair_id": f"audit-pair-{pair_index:03d}",
                "control_slot": slot,
                "event_id": f"event-{pair_index}",
                "app_id": f"app-{pair_index % 3}",
                "control_window_id": f"control-{pair_index}-{slot}",
                **{field: "" for field in CONTROL_AUDIT_FIELDS},
            }
            for pair_index in range(1, 21)
            for slot in range(1, 4)
        ]
    )
    template = tmp_path / "template.xlsx"
    _format_workbook(
        template,
        pair_frame=pair,
        control_frame=control,
        include_private=False,
    )
    input_dir = tmp_path / "results"
    private_dir = tmp_path / "private"
    output_dir = tmp_path / "summary"
    input_dir.mkdir()
    private_dir.mkdir()
    pd.DataFrame(
        [
            {
                "audit_pair_id": f"audit-pair-{index:03d}",
                "direction": "positive" if index % 2 else "negative",
            }
            for index in range(1, 21)
        ]
    ).to_csv(private_dir / "PRIVATE_PAIR_EVIDENCE.csv", index=False)
    for reviewer_index in range(1, 5):
        path = input_dir / f"reviewer_{reviewer_index}.xlsx"
        shutil.copy2(template, path)
        workbook = load_workbook(path)
        pair_sheet = workbook["组级核查"]
        pair_headers = {
            cell.value: cell.column for cell in pair_sheet[1]
        }
        for row in range(2, 22):
            pair_sheet.cell(
                row, pair_headers["all_three_controls_reviewed"], "是"
            )
            pair_sheet.cell(row, pair_headers["pair_conclusion"], "通过")
            pair_sheet.cell(
                row, pair_headers["auditor_id"], f"audit_rev_{reviewer_index:02d}"
            )
            pair_sheet.cell(row, pair_headers["audit_date"], "2026-07-24")
        control_sheet = workbook["逐对照核查"]
        control_headers = {
            cell.value: cell.column for cell in control_sheet[1]
        }
        for row in range(2, 62):
            for field in CONTROL_AUDIT_FIELDS[:7]:
                control_sheet.cell(row, control_headers[field], "通过")
            control_sheet.cell(
                row, control_headers["control_row_conclusion"], "通过"
            )
        workbook.save(path)
    return input_dir, private_dir, output_dir, template


def test_manual_audit_summary_accepts_four_complete_signed_reviews(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    result = summarize_manual_audit_results(
        input_dir,
        private_dir,
        output_dir,
        template_path=template,
    )
    assert result["completeness"]["control_ratings"] == 240
    assert result["completeness"]["pair_ratings"] == 80
    assert result["control_conclusions"] == {"通过": 240}
    assert result["pair_conclusions"] == {"通过": 80}
    assert result["unanimity"]["control_items_unanimous"] == 60
    assert result["unanimity"]["pair_items_unanimous"] == 20
    assert (
        result["unanimity"]["chance_corrected_agreement"]
        == "not_computed_constant_verdicts"
    )
    assert len(list((output_dir / "source_evidence").glob("*.xlsx"))) == 4


def test_manual_audit_summary_rejects_duplicate_workbooks(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    first = input_dir / "reviewer_1.xlsx"
    for index in range(2, 5):
        shutil.copy2(first, input_dir / f"reviewer_{index}.xlsx")
    with pytest.raises(ValueError, match="duplicate SHA-256"):
        summarize_manual_audit_results(
            input_dir,
            private_dir,
            output_dir,
            template_path=template,
        )


def test_manual_audit_summary_rejects_missing_signature(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    path = input_dir / "reviewer_4.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["组级核查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["auditor_id"]).value = None
    workbook.save(path)
    with pytest.raises(ValueError, match="blank auditor_id"):
        summarize_manual_audit_results(
            input_dir,
            private_dir,
            output_dir,
            template_path=template,
        )


def test_manual_audit_summary_rejects_duplicate_auditor_ids(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    path = input_dir / "reviewer_4.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["组级核查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, 22):
        sheet.cell(row, headers["auditor_id"], "audit_rev_01")
        sheet.cell(row, headers["audit_date"], "2026-07-25")
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicate auditor_id"):
        summarize_manual_audit_results(
            input_dir,
            private_dir,
            output_dir,
            template_path=template,
        )


def test_manual_audit_summary_rejects_changed_evidence(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    path = input_dir / "reviewer_4.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["逐对照核查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["event_id"], "tampered-event")
    workbook.save(path)
    with pytest.raises(ValueError, match="changed a frozen evidence value"):
        summarize_manual_audit_results(
            input_dir,
            private_dir,
            output_dir,
            template_path=template,
        )


def test_manual_audit_summary_rejects_partial_unfilled_review(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    path = input_dir / "reviewer_4.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["逐对照核查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["hour_distance_and_score_check"]).value = None
    workbook.save(path)
    with pytest.raises(
        ValueError,
        match="blank hour_distance_and_score_check values",
    ):
        summarize_manual_audit_results(
            input_dir,
            private_dir,
            output_dir,
            template_path=template,
        )


def test_manual_audit_summary_requires_issue_details_for_nonpass(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    path = input_dir / "reviewer_4.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["逐对照核查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["control_row_conclusion"], "需复核")
    workbook.save(path)
    with pytest.raises(ValueError, match="requires issue code and notes"):
        summarize_manual_audit_results(
            input_dir,
            private_dir,
            output_dir,
            template_path=template,
        )


def test_manual_audit_summary_accepts_valid_issue_code(
    tmp_path: Path,
) -> None:
    input_dir, private_dir, output_dir, template = _signed_audit_fixture(
        tmp_path
    )
    path = input_dir / "reviewer_4.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["逐对照核查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["control_row_conclusion"], "需复核")
    sheet.cell(2, headers["issue_codes"], "MATCH_SCORE")
    sheet.cell(
        2,
        headers["auditor_notes"],
        "independent recomputation differed",
    )
    workbook.save(path)
    result = summarize_manual_audit_results(
        input_dir,
        private_dir,
        output_dir,
        template_path=template,
    )
    assert result["control_conclusions"]["需复核"] == 1
    assert result["issue_codes"]["MATCH_SCORE"] == 1
