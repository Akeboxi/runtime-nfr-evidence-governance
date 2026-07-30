"""Build the Word submission draft from the auditable Markdown manuscript."""

from __future__ import annotations

import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE_PATTERN = re.compile(r"(\\\(.+?\\\)|\*\*.+?\*\*|`.+?`)")
IMAGE_PATTERN = re.compile(r"^!\[(.*?)\]\((.*?)\)$")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+)$")

SUBSCRIPT_TRANSLATION = str.maketrans(
    {
        "0": "₀",
        "1": "₁",
        "2": "₂",
        "3": "₃",
        "4": "₄",
        "5": "₅",
        "6": "₆",
        "7": "₇",
        "8": "₈",
        "9": "₉",
        ".": ".",
        "e": "ₑ",
        "h": "ₕ",
        "t": "ₜ",
        "v": "ᵥ",
    }
)


def _set_run_font(run, name: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _normalize_latex_math(latex: str) -> str:
    """Render the manuscript's small LaTeX subset as readable Unicode math."""
    text = latex.strip().rstrip(",")
    text = text.replace(
        r"\frac{Q_{0.99}(L)}{\max(Q_{0.50}(L),\epsilon)}",
        "Q₀.₉₉(L) / max(Q₀.₅₀(L), ε)",
    )
    text = text.replace(r"\frac{n_v}{n_e}", "nᵥ / nₑ")
    text = text.replace(r"\epsilon", "ε").replace(r"\max", "max")

    def replace_subscript(match: re.Match[str]) -> str:
        value = match.group(1) or match.group(2)
        return value.translate(SUBSCRIPT_TRANSLATION)

    text = re.sub(r"_\{([^{}]+)\}|_([A-Za-z0-9.]+)", replace_subscript, text)
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


def _add_math_run(paragraph, latex: str, *, size: float = 10.5) -> None:
    run = paragraph.add_run(_normalize_latex_math(latex))
    _set_run_font(run, "Cambria Math", size)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(
    cell, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def _add_inline(paragraph, text: str, *, default_size: float = 10.5) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            _set_run_font(run, "宋体", default_size)
        token = match.group(0)
        if token.startswith(r"\("):
            _add_math_run(paragraph, token[2:-2], size=default_size)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, "宋体", default_size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Consolas", max(default_size - 1, 8))
            run.font.color.rgb = RGBColor(31, 78, 121)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        _set_run_font(run, "宋体", default_size)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    _set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index == 0:
                _shade(cell, "D9EAF7")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value.strip())
            _set_run_font(run, "宋体", 8.5, bold=(row_index == 0))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    # Geometry measured from the official JOS "排版样例2025年版".
    section.page_width = Cm(18.4)
    section.page_height = Cm(26.0)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    style_specs = {
        "Title": ("黑体", 18, True, RGBColor(0, 0, 0)),
        "Heading 1": ("黑体", 14, True, RGBColor(31, 78, 121)),
        "Heading 2": ("黑体", 12, True, RGBColor(31, 78, 121)),
        "Heading 3": ("黑体", 10.5, True, RGBColor(31, 78, 121)),
    }
    for style_name, (font_name, size, bold, color) in style_specs.items():
        style = document.styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.first_line_indent = Cm(0)

    if "Caption" not in document.styles:
        document.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = document.styles["Caption"]
    caption.font.name = "宋体"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.bold = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.first_line_indent = Cm(0)

    header = section.header.paragraphs[0]
    header.text = "《软件学报》“智能化需求工程”专刊投稿稿"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(header.runs[0], "宋体", 8)
    header.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_field(footer)
    for run in footer.runs:
        _set_run_font(run, "Times New Roman", 8)


def _clean_markdown_cell(value: str) -> str:
    return value.strip().replace("**", "").replace("`", "")


def build_docx(markdown_path: Path, output_path: Path, figure_width_mm: float = 175.0) -> None:
    text = markdown_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    if lines and lines[0].strip() == "---":
        closing = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[closing + 1 :]

    document = Document()
    _configure_document(document)
    document.core_properties.title = "缺少显式 SLO 的云原生运行时非功能需求操作化"
    document.core_properties.subject = "《软件学报》“智能化需求工程”专刊投稿稿"
    document.core_properties.keywords = "NFR; SLO; runtime requirements; evidence governance"

    index = 0
    in_references = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue

        heading_match = HEADING_PATTERN.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2).strip()
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(10)
                run = paragraph.add_run(heading_text)
                _set_run_font(run, "黑体", 18, bold=True)
            else:
                style_level = min(level - 1, 3)
                document.add_paragraph(heading_text, style=f"Heading {style_level}")
            in_references = heading_text == "参考文献"
            index += 1
            continue

        if (
            line.startswith("**作者：")
            or line.startswith("**单位：")
            or line.startswith("**英文署名：")
            or line.startswith("**英文单位：")
            or line.startswith("**第一作者邮箱：")
            or line.startswith("**通信作者：")
        ):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(1)
            _add_inline(paragraph, line, default_size=10)
            index += 1
            continue

        image_match = IMAGE_PATTERN.match(line)
        if image_match:
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            section = document.sections[0]
            text_width_cm = (
                section.page_width - section.left_margin - section.right_margin
            ) / 360000
            overhang_cm = max(0.0, (figure_width_mm / 10.0 - text_width_cm) / 2.0)
            paragraph.paragraph_format.left_indent = Cm(-overhang_cm)
            paragraph.paragraph_format.right_indent = Cm(-overhang_cm)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run()
            inline_shape = run.add_picture(
                str(image_path), width=Cm(figure_width_mm / 10.0)
            )
            image_alt = image_match.group(1).strip() or image_path.stem
            inline_shape._inline.docPr.set("descr", image_alt)
            inline_shape._inline.docPr.set("title", image_alt)
            index += 1
            continue

        if (
            line.startswith("|")
            and index + 1 < len(lines)
            and lines[index + 1].strip().startswith("|")
        ):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [
                [_clean_markdown_cell(value) for value in table_line.strip("|").split("|")]
                for table_line in table_lines
            ]
            if len(parsed) >= 2 and all(set(value) <= {"-", ":", " "} for value in parsed[1]):
                parsed.pop(1)
            _add_table(document, parsed)
            continue

        if line.startswith("\\["):
            equation_lines = [line]
            index += 1
            while index < len(lines):
                equation_lines.append(lines[index].strip())
                if lines[index].strip().endswith("\\]"):
                    index += 1
                    break
                index += 1
            equation = " ".join(equation_lines).replace("\\[", "").replace("\\]", "").strip()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            _add_math_run(paragraph, equation)
            continue

        if re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0.74)
            _add_inline(paragraph, re.sub(r"^[-*]\s+", "", line))
            index += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0.74)
            _add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
            index += 1
            continue

        paragraph = document.add_paragraph()
        if line.startswith("**图 ") or line.startswith("**表 "):
            paragraph.style = document.styles["Caption"]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_together = True
        elif line.startswith("**关键词：") or line.startswith("**Keywords:**"):
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif in_references:
            paragraph.paragraph_format.first_line_indent = Cm(-0.7)
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(2)
        _add_inline(paragraph, line, default_size=9 if in_references else 10.5)
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default="docs/manuscript/runtime_nfr/RUNTIME_NFR_MANUSCRIPT.md",
    )
    parser.add_argument(
        "--output",
        default="docs/manuscript/runtime_nfr/Runtime_NFR_JOS_submission_draft_v2_175mm.docx",
    )
    parser.add_argument(
        "--figure-width-mm",
        type=float,
        default=175.0,
        help="Rendered figure width in millimetres (default: 175).",
    )
    args = parser.parse_args()
    build_docx(
        Path(args.input).resolve(),
        Path(args.output).resolve(),
        figure_width_mm=args.figure_width_mm,
    )


if __name__ == "__main__":
    main()
