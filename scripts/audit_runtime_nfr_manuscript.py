"""Audit key manuscript claims against frozen Runtime NFR paper sources."""

from __future__ import annotations

import argparse
from hashlib import sha256
import json
from pathlib import Path
from typing import Any

def _load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def audit(
    manuscript_path: Path,
    docx_path: Path,
    paper_tables_path: Path,
    cases_path: Path,
    output_path: Path,
    *,
    check_docx: bool = True,
) -> dict[str, Any]:
    manuscript = manuscript_path.read_text(encoding="utf-8")
    word_text = _docx_text(docx_path) if check_docx else ""
    tables = _load(paper_tables_path)
    cases = _load(cases_path)["cases"]
    governance = tables["tables"]["rq1_governance_distribution"]
    round2 = {
        row["dimension"]: row
        for row in tables["tables"]["round2_expert_review"]
    }
    prediction = {
        row["model"]: row
        for row in tables["tables"]["prediction_ablation"]
    }
    manual = tables["tables"]["manual_audit"]
    external = tables["tables"]["external_validation"]

    required_claims = [
        "8,668 张候选边界",
        f"{governance[3]['cards']:,} 张（{governance[3]['proportion']:.1%}）",
        f"{round2['threshold_plausibility']['median']:.0f}/5",
        f"只有 {round2['threshold_plausibility']['card_median_at_least_4']}/48 张卡",
        "48/48 张卡的治理适当性中位数均不低于 4",
        f"{manual['control_ratings']}/{manual['control_ratings']} 个逐对照判断",
        f"{manual['pair_ratings']}/{manual['pair_ratings']} 个组级判断通过",
        f"{external['cards']:,}/{external['cards']:,} 张卡路由至",
        f"PR-AUC 为 {prediction['app_logistic']['pr_auc']:.4f}",
        f"PR-AUC 最高（{prediction['app_hgb']['pr_auc']:.4f}）",
        (
            f"冻结边界案例的 \\(R_t={cases['candidate_for_stakeholder_review']['latency_tail_ratio']:.6f}\\)"
        ),
    ]
    markdown_checks = {
        claim: claim in manuscript
        for claim in required_claims
    }
    word_required = [
        "8,668 张候选边界",
        "48/48 张卡的治理适当性中位数均不低于 4",
        "240/240 个逐对照判断",
        "80/80 个组级判断通过",
        "PR-AUC 为 0.7878",
        "人工核查的验证职责限定为匹配实现与证据重建一致",
    ]
    word_checks = (
        {claim: claim in word_text for claim in word_required}
        if check_docx
        else {}
    )
    figure_checks = {
        f"figure_{index:02d}_": f"figure_{index:02d}_" in manuscript
        for index in range(1, 7)
    }
    boundary_checks = {
        "no_formal_slo_approval_claim": not any(
            forbidden in manuscript
            for forbidden in (
                "本文证明候选边界是正式 SLO",
                "专家确认阈值正确",
                "外部数据验证了阈值正确性",
            )
        ),
        "manual_audit_noncausal_boundary": (
            "因果效应、阈值正确性和正式 SLO 不属于该审计终点" in manuscript
        ),
        "zero_variance_no_kappa_alpha": (
            "不计算或解释 κ 或 α" in manuscript
        ),
        "external_validity_only": "`validity_only`" in manuscript,
        "topology_negative_result_bounded": (
            "其结论限定为当前静态、不完整表示" in manuscript
        ),
    }
    all_checks = {
        "markdown_claims": markdown_checks,
        "word_claims": word_checks,
        "figures": figure_checks,
        "claim_boundaries": boundary_checks,
    }
    values = [
        value
        for group in all_checks.values()
        for value in group.values()
    ]
    result = {
        "protocol": "runtime-nfr-v3-manuscript-numeric-audit/2",
        "sources": {
            "paper_tables": {
                "path": str(paper_tables_path.resolve()),
                "sha256": _hash(paper_tables_path),
            },
            "representative_cases": {
                "path": str(cases_path.resolve()),
                "sha256": _hash(cases_path),
            },
            "manuscript": {
                "path": str(manuscript_path.resolve()),
                "sha256": _hash(manuscript_path),
            },
            "word": (
                {
                    "path": str(docx_path.resolve()),
                    "sha256": _hash(docx_path),
                    "status": "checked",
                }
                if check_docx
                else {
                    "path": str(docx_path.resolve()),
                    "status": "deferred_until_content_freeze",
                }
            ),
        },
        "checks": all_checks,
        "pass": all(values),
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--manuscript",
        default="docs/manuscript/runtime_nfr/RUNTIME_NFR_MANUSCRIPT.md",
    )
    parser.add_argument(
        "--docx",
        default="docs/manuscript/runtime_nfr/Runtime_NFR_JOS_submission_draft.docx",
    )
    parser.add_argument(
        "--paper-tables",
        default="checkpoints/runtime_nfr_v3_academic/paper_tables/paper_tables.json",
    )
    parser.add_argument(
        "--cases",
        default="checkpoints/runtime_nfr_v3_academic/paper_artifacts/representative_cases.json",
    )
    parser.add_argument(
        "--output",
        default=(
            "checkpoints/runtime_nfr_v3_academic/paper_artifacts/"
            "manuscript_numeric_audit.json"
        ),
    )
    parser.add_argument(
        "--skip-docx",
        action="store_true",
        help="Run the Markdown/source audit only; defer DOCX checks until layout.",
    )
    args = parser.parse_args()
    result = audit(
        Path(args.manuscript),
        Path(args.docx),
        Path(args.paper_tables),
        Path(args.cases),
        Path(args.output),
        check_docx=not args.skip_docx,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["pass"]:
        raise RuntimeError("Runtime NFR manuscript numeric audit failed")


if __name__ == "__main__":
    main()
