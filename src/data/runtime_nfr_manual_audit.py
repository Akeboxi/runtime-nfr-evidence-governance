"""Blinded, evidence-complete manual audit package for matched controls."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timedelta
from hashlib import sha256
from pathlib import Path
import ast
import json
import math
import shutil
from typing import Any, Sequence

import pandas as pd

from .runtime_nfr_dataset import (
    RuntimeNFRSample,
    _app_alias,
    _read_segment_layout,
    _severity_for_boundaries,
    fit_nfr_boundary,
    load_app_sli_frame,
)
from .runtime_nfr_v2 import (
    _overlaps_any,
    _parse_timestamp,
    _raw_event_intervals,
    load_event_labels,
)


MANUAL_AUDIT_PROTOCOL = "runtime-nfr-v3-matched-control-manual-audit/2"
MANUAL_AUDIT_SUMMARY_PROTOCOL = "runtime-nfr-v3-manual-audit-summary/1"
AUDIT_DECISIONS = ("通过", "不通过", "无法判断")
ROW_CONCLUSIONS = ("通过", "需复核", "不通过")
ISSUE_CODES = (
    "TIME_WINDOW",
    "APP_MISMATCH",
    "REQUEST_VOLUME",
    "MATCH_SCORE",
    "CANDIDATE_RANK",
    "HISTORY_SAMPLE",
    "FUTURE_SAMPLE",
    "BUFFER_OVERLAP",
    "SOURCE_UNAVAILABLE",
    "OTHER",
)
PRIVATE_OUTCOME_FIELDS = (
    "event_breach",
    "control_breach",
    "control_breach_rate",
    "event_minus_control",
    "direction",
)
PRIVATE_COMPUTED_JUDGMENT_FIELDS = (
    "same_app_by_construction",
    "hour_distance_rule_pass",
    "candidate_rank_rule_pass",
    "history_sample_rule_pass",
    "future_sample_rule_pass",
    "buffer_nonoverlap_rule_pass",
)
CONTROL_AUDIT_FIELDS = (
    "same_app_check",
    "time_window_check",
    "request_volume_check",
    "hour_distance_and_score_check",
    "candidate_rank_check",
    "history_future_samples_check",
    "fault_buffer_check",
    "control_row_conclusion",
    "issue_codes",
    "auditor_notes",
)
PAIR_AUDIT_FIELDS = (
    "all_three_controls_reviewed",
    "pair_conclusion",
    "pair_issue_codes",
    "pair_notes",
    "auditor_id",
    "audit_date",
)


FIELD_DICTIONARY = [
    ("audit_pair_id", "盲化核查组 ID；一组对应一个事件—服务及三个对照。", "抽样程序"),
    ("control_slot", "同一核查组内的对照序号 1–3。", "匹配排序"),
    ("event_id", "内部事件 ID，仅用于定位原始材料，不包含结果方向。", "事件元数据"),
    ("segment_id", "事件所在内部数据段。", "事件元数据"),
    ("app_id", "事件和对照必须一致的服务 ID。", "事件/指标目录"),
    ("event_minute", "按分钟对齐的事件时点。", "事件元数据"),
    ("event_pre_window_start", "事件请求量匹配窗口起点，闭区间。", "event_minute-60 分钟"),
    ("event_pre_window_end_exclusive", "事件请求量匹配窗口终点，开区间。", "event_minute"),
    ("event_pre_request_volume", "事件前 60 分钟 request 求和。", "原始 App 指标"),
    ("event_pre_request_valid_samples", "事件前窗口中非空 request 分钟数。", "原始 App 指标"),
    ("control_window_id", "服务 ID 与对照起点生成的稳定哈希 ID。", "SHA-256 前 16 位"),
    ("control_start", "对照窗口起点。", "匹配程序"),
    ("control_observation_end_exclusive", "早期观察窗口终点。", "control_start+5 分钟"),
    ("control_future_end_exclusive", "未来窗口终点。", "control_start+20 分钟"),
    ("control_pre_window_start", "对照请求量匹配窗口起点。", "control_start-60 分钟"),
    ("control_pre_request_volume", "对照前 60 分钟 request 求和。", "原始 App 指标"),
    ("control_pre_request_valid_samples", "对照前窗口中非空 request 分钟数。", "原始 App 指标"),
    ("cyclic_hour_distance", "事件小时与对照小时的 24 小时循环距离。", "min(|h1-h2|,24-|h1-h2|)"),
    ("log_request_distance", "事件与对照请求量的 log1p 绝对距离。", "|log1p(control)-log1p(event)|"),
    ("matching_score", "冻结匹配分数，越小越优。", "小时距离+log 请求量距离"),
    ("candidate_rank", "通过预筛选的候选中按分数和时间排序的名次。", "匹配程序"),
    ("eligible_candidate_count", "该事件—服务通过预筛选的候选总数。", "匹配程序"),
    ("history_window_start", "对照边界历史窗口起点。", "control_start-1440 分钟"),
    ("history_valid_latency_samples", "剔除故障缓冲区后的有效延迟分钟数。", "原始 App 指标"),
    ("history_valid_error_samples", "剔除故障缓冲区后的有效错误率分钟数。", "原始 App 指标"),
    ("history_valid_samples_min", "延迟与错误率有效数的较小者；必须≥120。", "边界拟合"),
    ("history_rows_excluded_by_fault_buffer", "历史窗口内因故障缓冲区被剔除的已有数据行数。", "缓冲区审计"),
    ("future_valid_request_samples", "未来 15 分钟中有效 request 分钟数；必须≥8。", "未来窗口"),
    ("future_valid_latency_samples", "未来窗口中有效延迟分钟数。", "未来窗口"),
    ("future_valid_error_samples", "未来窗口中有效错误率分钟数。", "未来窗口"),
    ("selection_window_overlaps_fault_buffer", "从对照起点到未来窗口末端是否与任一故障缓冲区重叠；必须为 False。", "缓冲区审计"),
    ("selection_overlap_event_ids", "发生上述重叠时对应的事件 ID；合格对照应为空。", "缓冲区审计"),
    ("history_overlaps_fault_buffer", "24 小时历史窗口是否穿过故障缓冲区；允许为 True，但相应行必须已剔除。", "缓冲区审计"),
    ("history_overlap_event_ids", "历史窗口穿过的缓冲事件 ID。", "缓冲区审计"),
    ("minimum_distance_to_fault_buffer_minutes", "对照选择窗口距最近故障缓冲区的分钟距离。", "缓冲区审计"),
    ("event_metric_source_paths", "事件请求量证据的原始 CSV 路径。", "文件清单"),
    ("control_metric_source_paths", "对照历史、请求量和未来证据的原始 CSV 路径。", "文件清单"),
    ("source_segment_ids", "对照时间段实际涉及的数据段。", "文件清单"),
    ("*_check", "逐项填写“通过 / 不通过 / 无法判断”。", "核查人填写"),
    ("control_row_conclusion", "单个对照的结论：通过 / 需复核 / 不通过。", "核查人填写"),
    ("issue_codes", "问题代码；多个代码以英文分号连接。", "核查人填写"),
    ("auditor_notes", "原始证据缺失、重算差异或其他说明。", "核查人填写"),
    ("all_three_controls_reviewed", "该组 3 个对照是否均已完成逐行核查。", "核查人填写"),
    ("pair_conclusion", "事件—服务组的总体结论。", "核查人填写"),
    ("pair_issue_codes", "组级问题代码；多个代码以英文分号连接。", "核查人填写"),
    ("pair_notes", "组级说明和需要研究团队复核的问题。", "核查人填写"),
    ("auditor_id", "核查人匿名 ID。", "核查人填写"),
    ("audit_date", "核查完成日期，格式 YYYY-MM-DD。", "核查人填写"),
]


def _window_id(app_id: str, candidate: datetime) -> str:
    return sha256(f"{app_id}|{candidate.isoformat()}".encode("utf-8")).hexdigest()[:16]


def _direction(value: float) -> str:
    return "positive" if value > 0 else "negative" if value < 0 else "zero"


def _source_paths_for_window(
    sources: Sequence[dict[str, Any]],
    start: datetime,
    end: datetime,
) -> tuple[str, str]:
    paths: list[str] = []
    segments: list[str] = []
    start_stamp = pd.Timestamp(start)
    end_stamp = pd.Timestamp(end)
    for source in sources:
        frame = source["frame"]
        if not frame.empty and bool(
            ((frame.index >= start_stamp) & (frame.index < end_stamp)).any()
        ):
            paths.append(str(source["path"].resolve()))
            segments.append(str(source["segment_id"]))
    return ";".join(sorted(set(paths))), ";".join(sorted(set(segments)))


def _buffered_event_details(
    segment_root: Path,
    buffer_minutes: int,
) -> list[dict[str, Any]]:
    result = []
    for event_id, label in load_event_labels(segment_root).items():
        result.append(
            {
                "event_id": event_id,
                "start": _parse_timestamp(label["change_start_time"])
                - timedelta(minutes=buffer_minutes),
                "end": _parse_timestamp(label["change_end_time"])
                + timedelta(minutes=buffer_minutes),
            }
        )
    return result


def _overlap_ids(
    start: datetime,
    end: datetime,
    intervals: Sequence[dict[str, Any]],
) -> list[str]:
    return sorted(
        str(interval["event_id"])
        for interval in intervals
        if start < interval["end"] and end > interval["start"]
    )


def _minimum_buffer_distance_minutes(
    start: datetime,
    end: datetime,
    intervals: Sequence[dict[str, Any]],
) -> float:
    distances = []
    for interval in intervals:
        if start < interval["end"] and end > interval["start"]:
            return 0.0
        if end <= interval["start"]:
            distances.append((interval["start"] - end).total_seconds() / 60.0)
        else:
            distances.append((start - interval["end"]).total_seconds() / 60.0)
    return float(min(distances)) if distances else float("inf")


def _build_app_frames(
    segment_root: Path,
) -> tuple[dict[str, pd.DataFrame], dict[str, list[dict[str, Any]]]]:
    app_frames: dict[str, list[pd.DataFrame]] = defaultdict(list)
    sources: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for tenant_dir, meta, _, node_ids in _read_segment_layout(segment_root):
        segment_id = str(
            meta.get("segment_id") or tenant_dir.name.rsplit("_", 1)[-1]
        )
        for raw_app_id in node_ids["Vbiz"]:
            app_id = _app_alias(raw_app_id)
            path = (
                tenant_dir
                / "_staging"
                / "metrics"
                / "app"
                / f"{app_id}.csv"
            )
            if not path.exists():
                continue
            frame = load_app_sli_frame(path)
            if frame.empty:
                continue
            app_frames[app_id].append(frame)
            sources[app_id].append(
                {
                    "path": path,
                    "segment_id": segment_id,
                    "frame": frame,
                }
            )
    combined = {
        app_id: pd.concat(frames)
        .groupby(level=0, sort=True)
        .mean(numeric_only=True)
        for app_id, frames in app_frames.items()
    }
    return combined, sources


def build_private_matched_control_evidence(
    samples: Sequence[RuntimeNFRSample],
    sample_manifest_path: str | Path,
    *,
    segment_root: str | Path,
    history_minutes: int,
    observation_minutes: int,
    horizon_minutes: int,
    quantile: float,
    persistence: int,
    min_history_minutes: int,
    min_future_minutes: int,
    controls_per_pair: int,
    buffer_minutes: int,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Reconstruct selected controls with every matching and eligibility field."""

    selected = pd.read_csv(sample_manifest_path)
    if len(selected) != 20:
        raise ValueError("manual audit sample must contain exactly 20 event-App pairs")
    required = {
        "event_id",
        "segment_id",
        "app_id",
        "event_breach",
        "control_breach_rate",
        "event_minus_control",
        "control_window_ids",
    }
    missing = required - set(selected.columns)
    if missing:
        raise ValueError(f"manual audit sample missing columns: {sorted(missing)}")
    sample_by_event = {sample.metadata.event_id: sample for sample in samples}
    root = Path(segment_root)
    frames, sources = _build_app_frames(root)
    raw_intervals = _raw_event_intervals(root, buffer_minutes)
    detailed_intervals = _buffered_event_details(root, buffer_minutes)
    control_rows: list[dict[str, Any]] = []
    pair_rows: list[dict[str, Any]] = []

    for pair_index, selection in selected.reset_index(drop=True).iterrows():
        event_id = str(selection["event_id"])
        app_id = str(selection["app_id"])
        sample = sample_by_event.get(event_id)
        if sample is None:
            raise ValueError(f"selected event not found in dataset: {event_id}")
        frame = frames.get(app_id)
        if frame is None or frame.empty:
            raise ValueError(f"selected App has no telemetry: {app_id}")
        event_minute = sample.metadata.event_minute
        event_pre_start = event_minute - timedelta(minutes=60)
        event_history = frame.loc[
            (frame.index >= pd.Timestamp(event_pre_start))
            & (frame.index < pd.Timestamp(event_minute))
        ]
        event_pre_request = float(
            event_history["request"].fillna(0.0).sum()
        )
        event_pre_valid = int(event_history["request"].notna().sum())
        first = frame.index.min() + pd.Timedelta(
            minutes=max(min_history_minutes, 60)
        )
        last = frame.index.max() - pd.Timedelta(
            minutes=observation_minutes + horizon_minutes
        )
        candidates: list[dict[str, Any]] = []
        for timestamp in pd.date_range(first, last, freq="5min"):
            candidate = timestamp.to_pydatetime()
            hour_distance = min(
                abs(candidate.hour - event_minute.hour),
                24 - abs(candidate.hour - event_minute.hour),
            )
            if hour_distance > 1:
                continue
            selection_end = candidate + timedelta(
                minutes=observation_minutes + horizon_minutes
            )
            if _overlaps_any(candidate, selection_end, raw_intervals):
                continue
            pre_start = candidate - timedelta(minutes=60)
            pre = frame.loc[
                (frame.index >= pd.Timestamp(pre_start))
                & (frame.index < pd.Timestamp(candidate))
            ]
            pre_valid = int(pre["request"].notna().sum())
            if pre_valid < min(30, min_history_minutes):
                continue
            request = float(pre["request"].fillna(0.0).sum())
            log_request_distance = abs(
                math.log1p(request) - math.log1p(event_pre_request)
            )
            score = float(hour_distance + log_request_distance)
            candidates.append(
                {
                    "candidate": candidate,
                    "hour_distance": hour_distance,
                    "pre_start": pre_start,
                    "pre_request": request,
                    "pre_valid": pre_valid,
                    "log_request_distance": log_request_distance,
                    "score": score,
                }
            )
        candidates.sort(key=lambda row: (row["score"], row["candidate"]))
        for rank, candidate in enumerate(candidates, start=1):
            candidate["rank"] = rank
        chosen = candidates[:controls_per_pair]
        expected_ids = [
            str(value)
            for value in ast.literal_eval(str(selection["control_window_ids"]))
        ]
        chosen_ids = [
            _window_id(app_id, row["candidate"]) for row in chosen
        ]
        if chosen_ids != expected_ids:
            raise AssertionError(
                f"reconstructed controls differ for {event_id}/{app_id}: "
                f"{chosen_ids} != {expected_ids}"
            )
        audit_pair_id = f"audit-pair-{pair_index + 1:03d}"
        event_paths, _ = _source_paths_for_window(
            sources[app_id], event_pre_start, event_minute
        )
        pair_rows.append(
            {
                "audit_pair_id": audit_pair_id,
                "event_id": event_id,
                "event_date": sample.metadata.event_start.date().isoformat(),
                "segment_id": str(selection["segment_id"]),
                "app_id": app_id,
                "event_minute": event_minute.isoformat(),
                "event_pre_window_start": event_pre_start.isoformat(),
                "event_pre_window_end_exclusive": event_minute.isoformat(),
                "event_pre_request_volume": event_pre_request,
                "event_pre_request_valid_samples": event_pre_valid,
                "eligible_candidate_count": len(candidates),
                "selected_control_count": len(chosen),
                "selected_control_window_ids": ";".join(chosen_ids),
                "event_metric_source_paths": event_paths,
                "event_breach": float(selection["event_breach"]),
                "control_breach_rate": float(selection["control_breach_rate"]),
                "event_minus_control": float(selection["event_minus_control"]),
                "direction": _direction(float(selection["event_minus_control"])),
            }
        )
        for control_slot, candidate_row in enumerate(chosen, start=1):
            candidate = candidate_row["candidate"]
            history_start = candidate - timedelta(minutes=history_minutes)
            history_before = frame.loc[
                (frame.index >= pd.Timestamp(history_start))
                & (frame.index < pd.Timestamp(candidate))
            ]
            keep = pd.Series(True, index=history_before.index)
            for interval in raw_intervals:
                keep &= ~(
                    (history_before.index >= pd.Timestamp(interval.start))
                    & (history_before.index < pd.Timestamp(interval.end))
                )
            history = history_before.loc[keep]
            latency_boundary = fit_nfr_boundary(
                history["latency"],
                app_id=app_id,
                sli="latency",
                history_start=history_start,
                history_end=candidate,
                quantile=quantile,
            )
            error_boundary = fit_nfr_boundary(
                history["error_ratio"],
                app_id=app_id,
                sli="error_ratio",
                history_start=history_start,
                history_end=candidate,
                quantile=quantile,
            )
            future_start = candidate + timedelta(minutes=observation_minutes)
            future_end = future_start + timedelta(minutes=horizon_minutes)
            future_grid = pd.date_range(
                future_start,
                periods=horizon_minutes,
                freq="min",
            )
            future = frame.reindex(future_grid)
            control_breach, _, future_observed = _severity_for_boundaries(
                future,
                latency_boundary,
                error_boundary,
                persistence=persistence,
            )
            history_valid = min(
                latency_boundary.valid_samples,
                error_boundary.valid_samples,
            )
            if (
                history_valid < min_history_minutes
                or future_observed < min_future_minutes
            ):
                raise AssertionError(
                    f"frozen control no longer meets sample requirements: "
                    f"{event_id}/{app_id}/{candidate.isoformat()}"
                )
            selection_overlap_ids = _overlap_ids(
                candidate, future_end, detailed_intervals
            )
            history_overlap_ids = _overlap_ids(
                history_start, candidate, detailed_intervals
            )
            control_paths, control_segments = _source_paths_for_window(
                sources[app_id], history_start, future_end
            )
            control_rows.append(
                {
                    "audit_pair_id": audit_pair_id,
                    "control_slot": control_slot,
                    "event_id": event_id,
                    "segment_id": str(selection["segment_id"]),
                    "app_id": app_id,
                    "event_minute": event_minute.isoformat(),
                    "event_pre_window_start": event_pre_start.isoformat(),
                    "event_pre_window_end_exclusive": event_minute.isoformat(),
                    "event_pre_request_volume": event_pre_request,
                    "event_pre_request_valid_samples": event_pre_valid,
                    "control_window_id": _window_id(app_id, candidate),
                    "control_start": candidate.isoformat(),
                    "control_observation_end_exclusive": future_start.isoformat(),
                    "control_future_end_exclusive": future_end.isoformat(),
                    "control_pre_window_start": candidate_row[
                        "pre_start"
                    ].isoformat(),
                    "control_pre_window_end_exclusive": candidate.isoformat(),
                    "control_pre_request_volume": candidate_row["pre_request"],
                    "control_pre_request_valid_samples": candidate_row[
                        "pre_valid"
                    ],
                    "cyclic_hour_distance": candidate_row["hour_distance"],
                    "log_request_distance": candidate_row[
                        "log_request_distance"
                    ],
                    "matching_score": candidate_row["score"],
                    "candidate_rank": candidate_row["rank"],
                    "eligible_candidate_count": len(candidates),
                    "history_window_start": history_start.isoformat(),
                    "history_window_end_exclusive": candidate.isoformat(),
                    "history_rows_before_buffer_exclusion": len(history_before),
                    "history_rows_excluded_by_fault_buffer": int(
                        (~keep).sum()
                    ),
                    "history_valid_request_samples": int(
                        history["request"].notna().sum()
                    ),
                    "history_valid_latency_samples": int(
                        latency_boundary.valid_samples
                    ),
                    "history_valid_error_samples": int(
                        error_boundary.valid_samples
                    ),
                    "history_valid_samples_min": int(history_valid),
                    "future_window_start": future_start.isoformat(),
                    "future_window_end_exclusive": future_end.isoformat(),
                    "future_valid_request_samples": int(future_observed),
                    "future_valid_latency_samples": int(
                        future["latency"].notna().sum()
                    ),
                    "future_valid_error_samples": int(
                        future["error_ratio"].notna().sum()
                    ),
                    "selection_window_overlaps_fault_buffer": bool(
                        selection_overlap_ids
                    ),
                    "selection_overlap_event_ids": ";".join(
                        selection_overlap_ids
                    ),
                    "history_overlaps_fault_buffer": bool(
                        history_overlap_ids
                    ),
                    "history_overlap_event_ids": ";".join(
                        history_overlap_ids
                    ),
                    "minimum_distance_to_fault_buffer_minutes": (
                        _minimum_buffer_distance_minutes(
                            candidate, future_end, detailed_intervals
                        )
                    ),
                    "same_app_by_construction": True,
                    "hour_distance_rule_pass": (
                        candidate_row["hour_distance"] <= 1
                    ),
                    "candidate_rank_rule_pass": (
                        candidate_row["rank"] <= controls_per_pair
                    ),
                    "history_sample_rule_pass": (
                        history_valid >= min_history_minutes
                    ),
                    "future_sample_rule_pass": (
                        future_observed >= min_future_minutes
                    ),
                    "buffer_nonoverlap_rule_pass": not selection_overlap_ids,
                    "event_metric_source_paths": event_paths,
                    "control_metric_source_paths": control_paths,
                    "source_segment_ids": control_segments,
                    "event_breach": float(selection["event_breach"]),
                    "control_breach": bool(control_breach),
                    "control_breach_rate": float(
                        selection["control_breach_rate"]
                    ),
                    "event_minus_control": float(
                        selection["event_minus_control"]
                    ),
                    "direction": _direction(
                        float(selection["event_minus_control"])
                    ),
                }
            )
    private_controls = pd.DataFrame(control_rows)
    private_pairs = pd.DataFrame(pair_rows)
    if len(private_controls) != 60 or len(private_pairs) != 20:
        raise AssertionError("manual audit package must contain 20 pairs and 60 controls")
    if private_controls["selection_window_overlaps_fault_buffer"].any():
        raise AssertionError("selected control overlaps a buffered fault interval")
    return private_pairs, private_controls


def _blinded(frame: pd.DataFrame, audit_fields: Sequence[str]) -> pd.DataFrame:
    result = frame.drop(
        columns=[
            *PRIVATE_OUTCOME_FIELDS,
            *PRIVATE_COMPUTED_JUDGMENT_FIELDS,
        ],
        errors="ignore",
    ).copy()
    for field in audit_fields:
        result[field] = ""
    forbidden = set(PRIVATE_OUTCOME_FIELDS) & set(result.columns)
    if forbidden:
        raise AssertionError(f"blinded table retains outcomes: {sorted(forbidden)}")
    return result


def _write_instructions(path: Path) -> None:
    path.write_text(
        """# Runtime NFR 匹配对照人工核查说明

## 1. 核查目的

本核查只判断匹配对照的选择过程和证据是否正确，不评价事件是否造成越界，也不判断研究假设是否成立。核查单位为“事件—服务—对照窗口”：20 个事件—服务组，每组 3 个对照，共 60 行。

## 2. 盲化与独立性要求

1. 只使用 `MATCH_AUDIT_REVIEWER.xlsx` 或两个 `*_BLINDED.csv`。
2. 不打开 `private/` 目录，不查阅 `event_breach`、`control_breach`、`control_breach_rate`、`event_minus_control` 或 `direction`。
3. 不根据服务名、日期或指标高低猜测结果方向。
4. 独立完成，不与论文作者讨论某一行“应该通过还是失败”。遇到证据缺失时选择“无法判断”，不要补猜。
5. 核查结束前不与其他核查人交换判断。

## 3. 冻结匹配规则

- 同一服务精确匹配；
- 事件和对照均用起点前 60 分钟的 request 总量；
- 循环小时差 `min(|h1-h2|, 24-|h1-h2|)` 不得大于 1；
- 请求距离为 `|log1p(control_request)-log1p(event_request)|`；
- 匹配分数为“循环小时差＋请求距离”，按 `(matching_score, control_start)` 从小到大排序；
- 选择候选排名前 3 名；
- 候选选择窗口为 `[control_start, control_start+20min)`，不得与任何“故障开始前 30 分钟至故障结束后 30 分钟”的缓冲区重叠；
- 对照边界历史为 `[control_start-1440min, control_start)`；历史中落入故障缓冲区的已有行先剔除；
- 延迟和错误率历史有效样本数的较小者必须不少于 120；
- 未来窗口为 `[control_start+5min, control_start+20min)`，有效 request 分钟数必须不少于 8；
- 选择候选时不得使用未来越界结果。

## 4. 核查步骤

### A. 逐对照核查（60 行）

1. 根据 `event_metric_source_paths` 和 `control_metric_source_paths` 打开原始 CSV。
2. 核对事件时点、事件前 60 分钟和对照各时间窗边界。
3. 重新计算事件与对照 request 总量及非空分钟数。
4. 重新计算循环小时差、log 请求距离和匹配分数。
5. 检查 `candidate_rank` 是否为预筛选候选的真实排名 1–3。
6. 检查历史缓冲区剔除数、历史有效样本和未来有效样本。
7. 检查选择窗口是否与任何故障缓冲区重叠。历史窗口可以穿过缓冲区，但相应数据必须已经剔除。
8. 在七个 `*_check` 字段中填写“通过 / 不通过 / 无法判断”。
9. 填写 `control_row_conclusion`：
   - 全部通过：`通过`；
   - 存在无法判断但没有确定错误：`需复核`；
   - 任一关键规则不满足或证据与表格不一致：`不通过`。
10. 问题代码可多选，用分号连接，例如 `REQUEST_VOLUME;MATCH_SCORE`。

### B. 组级核查（20 行）

确认每组 3 行均已完成，再填写：

- `all_three_controls_reviewed`：通过/不通过/无法判断；
- `pair_conclusion`：通过/需复核/不通过；
- `pair_issue_codes` 和 `pair_notes`；
- 核查人匿名 ID 与日期。

## 5. 问题代码

`TIME_WINDOW`、`APP_MISMATCH`、`REQUEST_VOLUME`、`MATCH_SCORE`、`CANDIDATE_RANK`、`HISTORY_SAMPLE`、`FUTURE_SAMPLE`、`BUFFER_OVERLAP`、`SOURCE_UNAVAILABLE`、`OTHER`。

## 6. 完成标准与返还文件

- 60 行逐对照核查和 20 行组级核查均不得留空；
- “无法判断”必须说明缺少哪项证据；
- 不得修改灰色证据列，只填写黄色核查列；
- 返还填写后的 `MATCH_AUDIT_REVIEWER.xlsx`；
- 若使用 CSV，必须同时返还两个盲化 CSV，并保持 ID、行数和顺序不变；
- 核查人不得提交或接触 `private/` 目录。
""",
        encoding="utf-8",
    )


def _format_workbook(
    path: Path,
    *,
    pair_frame: pd.DataFrame,
    control_frame: pd.DataFrame,
    include_private: bool,
    private_pairs: pd.DataFrame | None = None,
    private_controls: pd.DataFrame | None = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    workbook.remove(workbook.active)
    instructions = workbook.create_sheet("说明")
    instruction_lines = [
        "Runtime NFR 匹配对照人工核查",
        "只核查匹配过程，不评价越界结果或论文假设。",
        "核查范围：20 个事件—服务组、每组 3 个对照、共 60 行。",
        "请先阅读同目录 MATCH_AUDIT_INSTRUCTIONS.md。",
        "黄色列由核查人填写；灰色列为冻结证据，不得修改。",
        "允许值：通过 / 不通过 / 无法判断；行结论：通过 / 需复核 / 不通过。",
        "不要打开或索取 private 目录，不要猜测结果方向。",
        "所有无法判断和不通过项必须填写问题代码与说明。",
    ]
    for row_index, line in enumerate(instruction_lines, start=1):
        instructions.cell(row_index, 1, line)
    instructions["A1"].font = Font(name="Arial", size=16, bold=True)
    instructions.column_dimensions["A"].width = 110
    instructions.freeze_panes = "A2"

    def add_frame(name: str, frame: pd.DataFrame, audit_columns: Sequence[str]) -> None:
        sheet = workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            sheet.append(list(row))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        audit_set = set(audit_columns)
        for column_index, column in enumerate(frame.columns, start=1):
            is_audit = column in audit_set
            fill = "FFF2CC" if is_audit else "E7E6E6"
            width = min(
                52,
                max(
                    12,
                    len(str(column)) + 2,
                    max(
                        (
                            len(str(value))
                            for value in frame[column].head(60)
                            if pd.notna(value)
                        ),
                        default=0,
                    )
                    + 2,
                ),
            )
            letter = sheet.cell(1, column_index).column_letter
            sheet.column_dimensions[letter].width = width
            for row_index in range(2, len(frame) + 2):
                cell = sheet.cell(row_index, column_index)
                cell.font = Font(name="Arial", size=9)
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        decision_validation = DataValidation(
            type="list",
            formula1='"通过,不通过,无法判断"',
            allow_blank=False,
        )
        conclusion_validation = DataValidation(
            type="list",
            formula1='"通过,需复核,不通过"',
            allow_blank=False,
        )
        sheet.add_data_validation(decision_validation)
        sheet.add_data_validation(conclusion_validation)
        for column_index, column in enumerate(frame.columns, start=1):
            letter = sheet.cell(1, column_index).column_letter
            cell_range = f"{letter}2:{letter}{len(frame) + 1}"
            if column.endswith("_check") or column in {
                "all_three_controls_reviewed"
            }:
                decision_validation.add(cell_range)
            if column in {"control_row_conclusion", "pair_conclusion"}:
                conclusion_validation.add(cell_range)
        if audit_columns:
            first_audit = frame.columns.get_loc(audit_columns[0]) + 1
            last_audit = frame.columns.get_loc(audit_columns[-1]) + 1
            start_letter = sheet.cell(1, first_audit).column_letter
            end_letter = sheet.cell(1, last_audit).column_letter
            sheet.conditional_formatting.add(
                f"{start_letter}2:{end_letter}{len(frame) + 1}",
                FormulaRule(
                    formula=[f'{start_letter}2="不通过"'],
                    fill=PatternFill("solid", fgColor="F4CCCC"),
                ),
            )

    add_frame("组级核查", pair_frame, PAIR_AUDIT_FIELDS)
    add_frame("逐对照核查", control_frame, CONTROL_AUDIT_FIELDS)
    dictionary = pd.DataFrame(
        FIELD_DICTIONARY, columns=["field", "description", "source_or_formula"]
    )
    add_frame("字段字典", dictionary, ())
    if include_private:
        if private_pairs is None or private_controls is None:
            raise ValueError("private workbook requires private evidence frames")
        add_frame("私有组级结果", private_pairs, ())
        add_frame("私有逐对照证据", private_controls, ())
        workbook["私有组级结果"].sheet_properties.tabColor = "C00000"
        workbook["私有逐对照证据"].sheet_properties.tabColor = "C00000"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)


def build_matched_control_audit_package(
    samples: Sequence[RuntimeNFRSample],
    sample_manifest_path: str | Path,
    output_dir: str | Path,
    *,
    segment_root: str | Path,
    history_minutes: int,
    observation_minutes: int,
    horizon_minutes: int,
    quantile: float,
    persistence: int,
    min_history_minutes: int,
    min_future_minutes: int,
    controls_per_pair: int = 3,
    buffer_minutes: int = 30,
) -> dict[str, Any]:
    output = Path(output_dir)
    reviewer_dir = output / "reviewer_package"
    private_dir = output / "private"
    reviewer_dir.mkdir(parents=True, exist_ok=True)
    private_dir.mkdir(parents=True, exist_ok=True)
    private_pairs, private_controls = build_private_matched_control_evidence(
        samples,
        sample_manifest_path,
        segment_root=segment_root,
        history_minutes=history_minutes,
        observation_minutes=observation_minutes,
        horizon_minutes=horizon_minutes,
        quantile=quantile,
        persistence=persistence,
        min_history_minutes=min_history_minutes,
        min_future_minutes=min_future_minutes,
        controls_per_pair=controls_per_pair,
        buffer_minutes=buffer_minutes,
    )
    blinded_pairs = _blinded(private_pairs, PAIR_AUDIT_FIELDS)
    blinded_controls = _blinded(private_controls, CONTROL_AUDIT_FIELDS)
    private_pair_csv = private_dir / "PRIVATE_PAIR_EVIDENCE.csv"
    private_control_csv = private_dir / "PRIVATE_CONTROL_EVIDENCE.csv"
    blinded_pair_csv = reviewer_dir / "MATCH_AUDIT_PAIR_SUMMARY_BLINDED.csv"
    blinded_control_csv = reviewer_dir / "MATCH_AUDIT_CONTROL_EVIDENCE_BLINDED.csv"
    private_pairs.to_csv(private_pair_csv, index=False, encoding="utf-8-sig")
    private_controls.to_csv(
        private_control_csv, index=False, encoding="utf-8-sig"
    )
    blinded_pairs.to_csv(blinded_pair_csv, index=False, encoding="utf-8-sig")
    blinded_controls.to_csv(
        blinded_control_csv, index=False, encoding="utf-8-sig"
    )
    instructions_path = reviewer_dir / "MATCH_AUDIT_INSTRUCTIONS.md"
    _write_instructions(instructions_path)
    reviewer_workbook = reviewer_dir / "MATCH_AUDIT_REVIEWER.xlsx"
    private_workbook = private_dir / "MATCH_AUDIT_PRIVATE.xlsx"
    _format_workbook(
        reviewer_workbook,
        pair_frame=blinded_pairs,
        control_frame=blinded_controls,
        include_private=False,
    )
    _format_workbook(
        private_workbook,
        pair_frame=blinded_pairs,
        control_frame=blinded_controls,
        include_private=True,
        private_pairs=private_pairs,
        private_controls=private_controls,
    )
    reviewer_headers = set(blinded_pairs.columns) | set(blinded_controls.columns)
    blinding_checks = {
        field: field not in reviewer_headers for field in PRIVATE_OUTCOME_FIELDS
    }
    computed_judgment_blinding_checks = {
        field: field not in reviewer_headers
        for field in PRIVATE_COMPUTED_JUDGMENT_FIELDS
    }
    files = [
        instructions_path,
        reviewer_workbook,
        private_workbook,
        private_pair_csv,
        private_control_csv,
        blinded_pair_csv,
        blinded_control_csv,
    ]
    manifest = {
        "protocol": MANUAL_AUDIT_PROTOCOL,
        "sample_manifest": str(Path(sample_manifest_path).resolve()),
        "sample_manifest_sha256": sha256(
            Path(sample_manifest_path).read_bytes()
        ).hexdigest(),
        "sample_selection_private": (
            "The frozen 20-pair sample was selected for case diversity across "
            "segment, application, and observed outcome direction. Outcome "
            "fields and the sampling direction are withheld from the auditor."
        ),
        "pairs": len(private_pairs),
        "controls": len(private_controls),
        "controls_per_pair": controls_per_pair,
        "blinding_checks": blinding_checks,
        "computed_judgment_blinding_checks": (
            computed_judgment_blinding_checks
        ),
        "reviewer_package_contains_private_outcomes": not all(
            blinding_checks.values()
        ),
        "matching_parameters": {
            "history_minutes": history_minutes,
            "request_match_window_minutes": 60,
            "observation_minutes": observation_minutes,
            "horizon_minutes": horizon_minutes,
            "quantile": quantile,
            "persistence": persistence,
            "min_history_minutes": min_history_minutes,
            "min_future_minutes": min_future_minutes,
            "buffer_minutes": buffer_minutes,
        },
        "files": [
            {
                "path": str(path.resolve()),
                "bytes": path.stat().st_size,
                "sha256": sha256(path.read_bytes()).hexdigest(),
            }
            for path in files
        ],
    }
    manifest_path = output / "audit_package_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {
        **manifest,
        "manifest": str(manifest_path.resolve()),
        "reviewer_workbook": str(reviewer_workbook.resolve()),
        "reviewer_instructions": str(instructions_path.resolve()),
        "private_workbook": str(private_workbook.resolve()),
        "private_control_evidence": str(private_control_csv.resolve()),
    }


def _file_record(path: Path) -> dict[str, Any]:
    return {
        "path": str(path.resolve()),
        "bytes": path.stat().st_size,
        "sha256": sha256(path.read_bytes()).hexdigest(),
    }


def _read_audit_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"{path.name} missing sheet: {sheet_name}")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{path.name}/{sheet_name} is empty")
    headers = [str(value) if value is not None else "" for value in rows[0]]
    if not all(headers):
        raise ValueError(f"{path.name}/{sheet_name} has blank headers")
    return pd.DataFrame(rows[1:], columns=headers)


def _normalized(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _validate_unchanged_evidence(
    result: pd.DataFrame,
    template: pd.DataFrame,
    audit_fields: Sequence[str],
    *,
    workbook_name: str,
    sheet_name: str,
) -> None:
    if list(result.columns) != list(template.columns):
        raise ValueError(f"{workbook_name}/{sheet_name} headers changed")
    if len(result) != len(template):
        raise ValueError(
            f"{workbook_name}/{sheet_name} row count changed: "
            f"{len(result)} != {len(template)}"
        )
    protected = [column for column in template.columns if column not in audit_fields]
    for row_index, (actual, expected) in enumerate(
        zip(
            result[protected].itertuples(index=False, name=None),
            template[protected].itertuples(index=False, name=None),
            strict=True,
        ),
        start=2,
    ):
        for column, actual_value, expected_value in zip(
            protected, actual, expected, strict=True
        ):
            if _normalized(actual_value) != _normalized(expected_value):
                raise ValueError(
                    f"{workbook_name}/{sheet_name}!{column}{row_index} "
                    "changed a frozen evidence value"
                )


def _split_issue_codes(value: Any) -> list[str]:
    if pd.isna(value) or not str(value).strip():
        return []
    return [
        part.strip()
        for part in str(value).replace("；", ";").split(";")
        if part.strip()
    ]


def _validate_issue_codes(
    frame: pd.DataFrame,
    *,
    code_field: str,
    conclusion_field: str,
    notes_field: str,
    workbook_name: str,
) -> Counter[str]:
    counts: Counter[str] = Counter()
    for row_index, row in frame.iterrows():
        codes = _split_issue_codes(row[code_field])
        invalid = sorted(set(codes) - set(ISSUE_CODES))
        if invalid:
            raise ValueError(
                f"{workbook_name} row {row_index + 2} has invalid issue codes: "
                f"{invalid}"
            )
        counts.update(codes)
        if row[conclusion_field] != "通过":
            if not codes or pd.isna(row[notes_field]) or not str(row[notes_field]).strip():
                raise ValueError(
                    f"{workbook_name} row {row_index + 2} requires issue code "
                    "and notes for a non-pass conclusion"
                )
    return counts


def summarize_manual_audit_results(
    input_dir: str | Path,
    private_dir: str | Path,
    output_dir: str | Path,
    *,
    template_path: str | Path,
    expected_reviewers: int = 4,
    expected_pairs: int = 20,
    expected_controls: int = 60,
) -> dict[str, Any]:
    """Validate, freeze, unblind, and summarize signed manual audit workbooks."""

    input_root = Path(input_dir)
    private_root = Path(private_dir)
    output_root = Path(output_dir)
    workbooks = sorted(input_root.glob("*.xlsx"))
    if len(workbooks) != expected_reviewers:
        raise ValueError(
            f"expected {expected_reviewers} reviewer workbooks, found {len(workbooks)}"
        )
    source_records = [_file_record(path) for path in workbooks]
    hashes = [record["sha256"] for record in source_records]
    if len(set(hashes)) != len(hashes):
        raise ValueError(
            "reviewer workbooks must be independently signed files; "
            "duplicate SHA-256 values were found"
        )

    template = Path(template_path)
    template_pairs = _read_audit_sheet(template, "组级核查")
    template_controls = _read_audit_sheet(template, "逐对照核查")
    reviewer_summaries: list[dict[str, Any]] = []
    pair_frames: list[pd.DataFrame] = []
    control_frames: list[pd.DataFrame] = []
    seen_auditors: set[str] = set()
    all_issue_counts: Counter[str] = Counter()

    for path in workbooks:
        pairs = _read_audit_sheet(path, "组级核查")
        controls = _read_audit_sheet(path, "逐对照核查")
        _validate_unchanged_evidence(
            pairs,
            template_pairs,
            PAIR_AUDIT_FIELDS,
            workbook_name=path.name,
            sheet_name="组级核查",
        )
        _validate_unchanged_evidence(
            controls,
            template_controls,
            CONTROL_AUDIT_FIELDS,
            workbook_name=path.name,
            sheet_name="逐对照核查",
        )
        if len(pairs) != expected_pairs or len(controls) != expected_controls:
            raise ValueError(
                f"{path.name} must contain {expected_pairs} pairs and "
                f"{expected_controls} controls"
            )

        required_control_fields = [
            *CONTROL_AUDIT_FIELDS[:8],
        ]
        for field in required_control_fields:
            if controls[field].isna().any() or (
                controls[field].astype(str).str.strip() == ""
            ).any():
                raise ValueError(f"{path.name} has blank {field} values")
        check_fields = CONTROL_AUDIT_FIELDS[:7]
        for field in check_fields:
            invalid = sorted(set(controls[field]) - set(AUDIT_DECISIONS))
            if invalid:
                raise ValueError(f"{path.name} has invalid {field}: {invalid}")
        invalid_rows = sorted(
            set(controls["control_row_conclusion"]) - set(ROW_CONCLUSIONS)
        )
        if invalid_rows:
            raise ValueError(
                f"{path.name} has invalid control conclusions: {invalid_rows}"
            )

        for field in ("all_three_controls_reviewed", "pair_conclusion", "auditor_id", "audit_date"):
            if pairs[field].isna().any() or (
                pairs[field].astype(str).str.strip() == ""
            ).any():
                raise ValueError(f"{path.name} has blank {field} values")
        invalid_reviewed = sorted(
            set(pairs["all_three_controls_reviewed"])
            - {"是", *AUDIT_DECISIONS}
        )
        if invalid_reviewed:
            raise ValueError(
                f"{path.name} has invalid all_three_controls_reviewed: "
                f"{invalid_reviewed}"
            )
        invalid_pairs = sorted(
            set(pairs["pair_conclusion"]) - set(ROW_CONCLUSIONS)
        )
        if invalid_pairs:
            raise ValueError(
                f"{path.name} has invalid pair conclusions: {invalid_pairs}"
            )
        auditor_ids = {
            str(value).strip() for value in pairs["auditor_id"] if str(value).strip()
        }
        if len(auditor_ids) != 1:
            raise ValueError(f"{path.name} must contain one consistent auditor_id")
        auditor_id = next(iter(auditor_ids))
        if auditor_id in seen_auditors:
            raise ValueError(f"duplicate auditor_id: {auditor_id}")
        seen_auditors.add(auditor_id)
        for value in pairs["audit_date"]:
            try:
                datetime.fromisoformat(str(value))
            except ValueError as error:
                raise ValueError(
                    f"{path.name} has invalid audit_date: {value}"
                ) from error

        issue_counts = _validate_issue_codes(
            controls,
            code_field="issue_codes",
            conclusion_field="control_row_conclusion",
            notes_field="auditor_notes",
            workbook_name=path.name,
        )
        issue_counts.update(
            _validate_issue_codes(
                pairs,
                code_field="pair_issue_codes",
                conclusion_field="pair_conclusion",
                notes_field="pair_notes",
                workbook_name=path.name,
            )
        )
        all_issue_counts.update(issue_counts)
        pairs = pairs.assign(auditor_id=auditor_id)
        controls = controls.assign(auditor_id=auditor_id)
        pair_frames.append(pairs)
        control_frames.append(controls)
        reviewer_summaries.append(
            {
                "auditor_id": auditor_id,
                "source_file": path.name,
                "controls": len(controls),
                "pairs": len(pairs),
                "control_conclusions": dict(
                    Counter(controls["control_row_conclusion"])
                ),
                "pair_conclusions": dict(Counter(pairs["pair_conclusion"])),
                "issue_codes": dict(issue_counts),
            }
        )

    all_pairs = pd.concat(pair_frames, ignore_index=True)
    all_controls = pd.concat(control_frames, ignore_index=True)
    pair_matrix = all_pairs.pivot(
        index="audit_pair_id", columns="auditor_id", values="pair_conclusion"
    )
    control_matrix = all_controls.pivot(
        index=["audit_pair_id", "control_slot"],
        columns="auditor_id",
        values="control_row_conclusion",
    )
    unanimous_pairs = int((pair_matrix.nunique(axis=1) == 1).sum())
    unanimous_controls = int((control_matrix.nunique(axis=1) == 1).sum())

    direction_summary: dict[str, Any] = {}
    private_pair_path = private_root / "PRIVATE_PAIR_EVIDENCE.csv"
    if private_pair_path.exists():
        private_pairs = pd.read_csv(private_pair_path)
        if {"audit_pair_id", "direction"}.issubset(private_pairs.columns):
            unblinded = all_pairs.merge(
                private_pairs[["audit_pair_id", "direction"]],
                on="audit_pair_id",
                how="left",
                validate="many_to_one",
            )
            direction_summary = {
                str(direction): {
                    "ratings": int(len(group)),
                    "conclusions": dict(Counter(group["pair_conclusion"])),
                }
                for direction, group in unblinded.groupby("direction", dropna=False)
            }

    output_root.mkdir(parents=True, exist_ok=True)
    frozen_source_dir = output_root / "source_evidence"
    frozen_source_dir.mkdir(parents=True, exist_ok=True)
    for path in workbooks:
        shutil.copy2(path, frozen_source_dir / path.name)
    frozen_records = [
        _file_record(frozen_source_dir / path.name) for path in workbooks
    ]
    summary = {
        "protocol": MANUAL_AUDIT_SUMMARY_PROTOCOL,
        "status": "complete",
        "source_files": frozen_records,
        "template": _file_record(template),
        "reviewers": reviewer_summaries,
        "completeness": {
            "reviewers": len(reviewer_summaries),
            "controls_per_reviewer": expected_controls,
            "pairs_per_reviewer": expected_pairs,
            "control_ratings": len(all_controls),
            "pair_ratings": len(all_pairs),
        },
        "control_conclusions": dict(
            Counter(all_controls["control_row_conclusion"])
        ),
        "pair_conclusions": dict(Counter(all_pairs["pair_conclusion"])),
        "unanimity": {
            "control_items_unanimous": unanimous_controls,
            "control_items_total": expected_controls,
            "pair_items_unanimous": unanimous_pairs,
            "pair_items_total": expected_pairs,
            "chance_corrected_agreement": (
                "not_computed_constant_verdicts"
                if all_controls["control_row_conclusion"].nunique() == 1
                and all_pairs["pair_conclusion"].nunique() == 1
                else "not_requested"
            ),
        },
        "issue_codes": dict(all_issue_counts),
        "unblinded_direction_summary": direction_summary,
        "claim_boundary": (
            "The audit supports consistency of matching implementation and "
            "evidence reconstruction; it does not establish causal effects, "
            "threshold correctness, or formal SLO approval."
        ),
    }
    summary_path = output_root / "manual_audit_summary.json"
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    hashes_path = output_root / "source_evidence_hashes.json"
    hashes_path.write_text(
        json.dumps(
            {
                "protocol": MANUAL_AUDIT_SUMMARY_PROTOCOL,
                "files": frozen_records,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    report_path = output_root / "MANUAL_AUDIT_REPORT.md"
    report_path.write_text(
        "\n".join(
            [
                "# Runtime NFR 匹配对照人工核查报告",
                "",
                f"- 核查人：{len(reviewer_summaries)}",
                f"- 逐对照判断：{len(all_controls)}",
                f"- 组级判断：{len(all_pairs)}",
                (
                    "- 逐对照结论："
                    + "，".join(
                        f"{key} {value}"
                        for key, value in summary["control_conclusions"].items()
                    )
                ),
                (
                    "- 组级结论："
                    + "，".join(
                        f"{key} {value}"
                        for key, value in summary["pair_conclusions"].items()
                    )
                ),
                (
                    f"- 一致核查项：逐对照 {unanimous_controls}/"
                    f"{expected_controls}，组级 {unanimous_pairs}/{expected_pairs}"
                ),
                "- 一致性系数：结论无变异时不计算，仅报告原始一致率。",
                "",
                "## 主张边界",
                "",
                "人工核查支持匹配实现与证据重建的一致性；不能据此推出因果效应、阈值正确性或正式 SLO 获批。",
                "",
            ]
        ),
        encoding="utf-8",
    )
    return {
        **summary,
        "summary_path": str(summary_path.resolve()),
        "report_path": str(report_path.resolve()),
        "source_hashes_path": str(hashes_path.resolve()),
    }
